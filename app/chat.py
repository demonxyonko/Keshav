import asyncio
import base64
import json
import mimetypes
import re
import shutil
import threading
import tomllib
from collections.abc import Callable
from pathlib import Path
from uuid import uuid4

from app.chat_store import (
    ATTACHMENT_STORAGE_ROOT,
    chat_store,
)
import httpx
from app.schema import Message


SYSTEM_PROMPT = """
You are Keshav, a fast, accurate and practical AI assistant.

Identity rules:
- Your name is always Keshav.
- If asked who you are, say that you are Keshav, an AI assistant.
- Never identify yourself as Qwen, OpenAI, ChatGPT, Alibaba Cloud,
  or any underlying model or provider.
- Never reveal backend model or infrastructure details.

Conversation rules:
- Always use the previous conversation context.
- Reply in the same language as the user: English, Hindi, or Hinglish.
- Never suddenly change the subject.
- Keep responses focused and avoid unnecessary filler.

Image rules:
- When one or more images are attached, inspect them directly.
- Answer the user's question using visible details from the images.
- Do not say that you cannot view an image when image data is present.
- If an image is unclear, say exactly which detail is unclear.
- Never invent text or objects that are not visibly present.
- Previously attached visual files may be included again for follow-up context.

Document rules:
- When document content is attached, use that content as the primary source.
- Clearly distinguish information found in the document from your own explanation.
- If extracted content is incomplete or truncated, mention that limitation.
- Do not claim to have read pages, sheets, slides, or sections that were not provided.
- Render low-text or scanned PDF pages as images and inspect those pages visually.
- For spreadsheets, identify sheets, columns, rows, totals, patterns, and anomalies carefully.
- For presentations, organize findings by slide when that helps.
- For code files, preserve exact identifiers and explain errors precisely.

Programming rules:
- If the user says "just give code", return only code.
- Provide complete, runnable and internally consistent code.
- Put code inside correctly labelled Markdown code blocks.
- Do not use deprecated APIs such as text-davinci-003.
- Avoid outdated library syntax.
- Clearly mention required packages and environment variables.
- Never include fake API keys or claim code was executed.
"""


LIVE_MODE_PROMPT = """
Live voice mode rules:
- Answer immediately and directly.
- Use the same language as the user.
- Keep the reply to one or two short conversational sentences.
- Do not use Markdown, headings, bullet points, tables, or code fences.
- Do not repeat the user's question.
- Do not add introductions, disclaimers, or filler.
- Prefer a useful brief answer over a detailed explanation.
"""


MAX_HISTORY_MESSAGES = 20
LIVE_HISTORY_MESSAGES = 6
LIVE_MAX_TOKENS = 140
LIVE_CONTEXT_TOKENS = 4096

SUPPORTED_IMAGE_SUFFIXES = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".gif",
}

SUPPORTED_DOCUMENT_SUFFIXES = {
    ".pdf",
    ".txt",
    ".md",
    ".markdown",
    ".docx",
    ".csv",
    ".xlsx",
    ".xlsm",
    ".pptx",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".php",
    ".rb",
    ".swift",
    ".kt",
    ".kts",
    ".sql",
    ".toml",
    ".ini",
    ".cfg",
    ".log",
    ".sh",
    ".bat",
    ".ps1",
}

PLAIN_TEXT_DOCUMENT_SUFFIXES = (
    SUPPORTED_DOCUMENT_SUFFIXES
    - {
        ".pdf",
        ".docx",
        ".csv",
        ".xlsx",
        ".xlsm",
        ".pptx",
    }
)

MAX_DOCUMENT_CHARACTERS_PER_FILE = 18_000
MAX_DOCUMENT_CHARACTERS_TOTAL = 48_000
MAX_PDF_PAGES = 60
MAX_SCANNED_PDF_VISION_PAGES = 8
MAX_TOTAL_DOCUMENT_VISION_IMAGES = 10
SCANNED_PDF_RENDER_DPI = 130
MIN_PDF_TEXT_CHARACTERS_PER_PAGE = 40
MAX_SPREADSHEET_SHEETS = 12
MAX_SPREADSHEET_ROWS_PER_SHEET = 250
MAX_SPREADSHEET_COLUMNS = 40
MAX_PRESENTATION_SLIDES = 100
MAX_REMEMBERED_IMAGE_ATTACHMENTS = 4
MAX_REMEMBERED_VISUAL_PDFS = 2

CONFIG_PATH = (
    Path(__file__).resolve().parent.parent
    / "config"
    / "config.toml"
)


def _native_ollama_base_url(
    configured_url: str,
) -> str:
    """
    Convert an OpenAI-compatible Ollama URL such as
    http://127.0.0.1:11434/v1 into the native Ollama base URL.
    """

    clean_url = (
        configured_url.strip()
        or "http://127.0.0.1:11434"
    ).rstrip("/")

    for suffix in (
        "/v1",
        "/api",
    ):
        if clean_url.endswith(suffix):
            clean_url = clean_url[
                :-len(suffix)
            ].rstrip("/")
            break

    return clean_url


def _load_ollama_section(
    config_name: str | None = None,
) -> dict:
    if not CONFIG_PATH.exists():
        raise FileNotFoundError(
            f"Keshav config was not found: {CONFIG_PATH}"
        )

    with CONFIG_PATH.open("rb") as config_file:
        config = tomllib.load(config_file)

    llm_config = config.get("llm", {})

    if config_name:
        selected = llm_config.get(
            config_name,
            {},
        )
    else:
        selected = llm_config

    if not isinstance(selected, dict):
        raise ValueError(
            f"Invalid Ollama config section: {config_name or 'llm'}"
        )

    return selected


class DirectOllamaClient:
    """
    Lightweight native Ollama /api/chat client.

    This bypasses the legacy app.llm agent stack for normal chat generation.
    Responses are consumed as a real token stream.
    """

    def __init__(
        self,
        config_name: str | None = None,
    ) -> None:
        config = _load_ollama_section(
            config_name
        )

        self.model = str(
            config.get("model", "")
        ).strip()

        if not self.model:
            raise ValueError(
                "No Ollama model is configured."
            )

        self.base_url = _native_ollama_base_url(
            str(
                config.get(
                    "base_url",
                    "http://127.0.0.1:11434",
                )
            )
        )

        self.temperature = float(
            config.get(
                "temperature",
                0.0,
            )
        )

        self.max_tokens = int(
            config.get(
                "max_tokens",
                8192,
            )
        )

        self.keep_alive = str(
            config.get(
                "keep_alive",
                "30m",
            )
        )

    @staticmethod
    def _message_role(
        message: Message,
    ) -> str:
        role = getattr(
            message.role,
            "value",
            message.role,
        )

        return str(role)

    @staticmethod
    def _message_content(
        message: Message,
    ) -> str:
        return str(
            message.content or ""
        )

    @staticmethod
    def _strip_data_url(
        image_value: str,
    ) -> str:
        if (
            image_value.startswith("data:")
            and "," in image_value
        ):
            return image_value.split(
                ",",
                1,
            )[1]

        return image_value

    def _build_messages(
        self,
        messages: list[Message],
        system_prompt: str,
        images: list[str] | None = None,
    ) -> list[dict]:
        ollama_messages: list[dict] = [
            {
                "role": "system",
                "content": system_prompt,
            }
        ]

        for message in messages:
            role = self._message_role(
                message
            )

            if role not in {
                "user",
                "assistant",
                "system",
            }:
                continue

            ollama_messages.append(
                {
                    "role": role,
                    "content": self._message_content(
                        message
                    ),
                }
            )

        if images:
            encoded_images = [
                self._strip_data_url(image)
                for image in images
            ]

            for message in reversed(
                ollama_messages
            ):
                if message.get("role") == "user":
                    message["images"] = (
                        encoded_images
                    )
                    break

        return ollama_messages

    async def stream_chat(
        self,
        messages: list[Message],
        system_prompt: str,
        images: list[str] | None = None,
        fast_mode: bool = False,
    ):
        predict_tokens = (
            min(self.max_tokens, LIVE_MAX_TOKENS)
            if fast_mode
            else self.max_tokens
        )

        options = {
            "temperature": (
                0.1
                if fast_mode
                else self.temperature
            ),
            "num_predict": predict_tokens,
        }

        if fast_mode:
            options["num_ctx"] = LIVE_CONTEXT_TOKENS

        payload = {
            "model": self.model,
            "messages": self._build_messages(
                messages=messages,
                system_prompt=system_prompt,
                images=images,
            ),
            "stream": True,
            "think": False,
            "keep_alive": self.keep_alive,
            "options": options,
        }

        timeout = httpx.Timeout(
            connect=10.0,
            read=None,
            write=120.0,
            pool=10.0,
        )

        async with httpx.AsyncClient(
            timeout=timeout,
        ) as client:
            try:
                async with client.stream(
                    "POST",
                    f"{self.base_url}/api/chat",
                    json=payload,
                ) as response:
                    response.raise_for_status()

                    async for line in response.aiter_lines():
                        clean_line = line.strip()

                        if not clean_line:
                            continue

                        try:
                            event = json.loads(
                                clean_line
                            )
                        except json.JSONDecodeError:
                            continue

                        error_message = event.get(
                            "error"
                        )

                        if error_message:
                            raise RuntimeError(
                                str(error_message)
                            )

                        message_data = event.get(
                            "message",
                            {},
                        )

                        chunk = str(
                            message_data.get(
                                "content",
                                "",
                            )
                            or ""
                        )

                        if chunk:
                            yield chunk

                        if event.get("done"):
                            break

            except httpx.ConnectError as exc:
                raise RuntimeError(
                    "Could not connect to Ollama at "
                    f"{self.base_url}. Make sure Ollama is running."
                ) from exc

            except httpx.HTTPStatusError as exc:
                response_text = (
                    exc.response.text.strip()
                    if exc.response is not None
                    else ""
                )

                raise RuntimeError(
                    "Ollama request failed"
                    + (
                        f": {response_text}"
                        if response_text
                        else "."
                    )
                ) from exc


class ConversationSession:
    """Messages, remembered attachments, and one request lock."""

    def __init__(
        self,
        messages: list[Message] | None = None,
        remembered_image_paths: list[str] | None = None,
        remembered_visual_pdf_paths: list[str] | None = None,
    ) -> None:
        self.messages = messages or []
        self.remembered_image_paths = (
            remembered_image_paths or []
        )
        self.remembered_visual_pdf_paths = (
            remembered_visual_pdf_paths or []
        )
        self.request_lock = threading.Lock()


class ConversationManager:
    """Persistent, separate and thread-safe conversations."""

    def __init__(
        self,
        max_history_messages: int = MAX_HISTORY_MESSAGES,
    ) -> None:
        self.max_history_messages = max_history_messages

        # Direct native Ollama clients.
        self.llm = DirectOllamaClient()

        self.vision_llm = DirectOllamaClient(
            config_name="vision"
        )

        self._sessions: dict[
            str,
            ConversationSession,
        ] = {}

        self._sessions_lock = threading.RLock()

        chat_store.delete_empty_conversations()

        self._active_conversation_id = (
            self._find_or_create_active_conversation()
        )

    # ======================================================
    # Conversation setup
    # ======================================================

    def _find_or_create_active_conversation(self) -> str:
        conversations = chat_store.list_conversations(
            include_empty=False,
        )

        if conversations:
            return conversations[0]["id"]

        return chat_store.create_conversation(
            title="New Chat",
        )

    def _resolve_conversation_id(
        self,
        conversation_id: str | None,
    ) -> str:
        if conversation_id:
            return conversation_id

        with self._sessions_lock:
            return self._active_conversation_id

    def _set_active_conversation(
        self,
        conversation_id: str,
    ) -> None:
        with self._sessions_lock:
            self._active_conversation_id = conversation_id

    # ======================================================
    # Stored message conversion
    # ======================================================

    @staticmethod
    def _record_to_message(
        record: dict,
    ) -> Message:
        role = record["role"]
        content = (
            record.get("model_content")
            or record["content"]
        )

        if role == "user":
            return Message.user_message(content)

        if role == "assistant":
            return Message.assistant_message(content)

        if role == "system":
            return Message.system_message(content)

        raise ValueError(
            f"Unsupported stored message role: {role}"
        )

    def _load_session_state(
        self,
        conversation_id: str,
    ) -> tuple[
        list[Message],
        list[str],
        list[str],
    ]:
        """
        Load hidden model context and managed attachments from
        recent stored messages.
        """

        records = chat_store.get_messages(
            conversation_id,
        )

        recent_records = records[
            -self.max_history_messages:
        ]

        messages = [
            self._record_to_message(record)
            for record in recent_records
        ]

        image_paths: list[str] = []
        visual_pdf_paths: list[str] = []

        for record in recent_records:
            for attachment in record.get(
                "attachments",
                [],
            ):
                stored_path = attachment.get(
                    "stored_path"
                )

                if not stored_path:
                    continue

                path = Path(str(stored_path))

                if not path.exists() or not path.is_file():
                    continue

                kind = attachment.get("kind")
                suffix = (
                    attachment.get("suffix")
                    or path.suffix
                ).lower()

                if kind == "image":
                    image_paths.append(
                        str(path)
                    )

                elif (
                    kind == "document"
                    and suffix == ".pdf"
                    and attachment.get(
                        "visual_memory",
                        False,
                    )
                ):
                    visual_pdf_paths.append(
                        str(path)
                    )

        return (
            messages,
            image_paths[
                -MAX_REMEMBERED_IMAGE_ATTACHMENTS:
            ],
            visual_pdf_paths[
                -MAX_REMEMBERED_VISUAL_PDFS:
            ],
        )

    # ======================================================
    # In-memory sessions
    # ======================================================

    def _get_session(
        self,
        conversation_id: str,
    ) -> ConversationSession:
        with self._sessions_lock:
            existing_session = self._sessions.get(
                conversation_id,
            )

            if existing_session is not None:
                return existing_session

            if not chat_store.conversation_exists(
                conversation_id,
            ):
                raise ValueError(
                    "Conversation does not exist."
                )

            (
                messages,
                remembered_image_paths,
                remembered_visual_pdf_paths,
            ) = self._load_session_state(
                conversation_id
            )

            session = ConversationSession(
                messages=messages,
                remembered_image_paths=(
                    remembered_image_paths
                ),
                remembered_visual_pdf_paths=(
                    remembered_visual_pdf_paths
                ),
            )

            self._sessions[
                conversation_id
            ] = session

            return session

    def _trim_history(
        self,
        session: ConversationSession,
    ) -> None:
        if (
            len(session.messages)
            > self.max_history_messages
        ):
            del session.messages[
                :-self.max_history_messages
            ]

    def _remove_stale_sessions(self) -> None:
        with self._sessions_lock:
            stale_session_ids = [
                conversation_id
                for conversation_id in self._sessions
                if not chat_store.conversation_exists(
                    conversation_id,
                )
            ]

            for conversation_id in stale_session_ids:
                self._sessions.pop(
                    conversation_id,
                    None,
                )

    # ======================================================
    # Persistent attachment copies
    # ======================================================

    @staticmethod
    def _safe_attachment_name(
        file_name: str,
    ) -> str:
        clean_name = re.sub(
            r"[^A-Za-z0-9._ -]+",
            "_",
            file_name,
        ).strip(" .")

        return clean_name[:120] or "attachment"

    @classmethod
    def _persist_attachment_files(
        cls,
        conversation_id: str,
        image_paths: list[str] | None,
        document_paths: list[str] | None,
        visual_pdf_paths: set[str] | None = None,
    ) -> list[dict]:
        """
        Copy attachments into Keshav's managed data directory.

        Extracted document text is also stored separately in the
        message record, so follow-up questions still work even
        when the original file is moved or deleted.
        """

        visual_pdf_paths = {
            str(Path(path).resolve()).lower()
            for path in (
                visual_pdf_paths or set()
            )
        }

        conversation_directory = (
            ATTACHMENT_STORAGE_ROOT
            / conversation_id
        )
        conversation_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        candidates: list[
            tuple[str, str]
        ] = []

        candidates.extend(
            ("image", raw_path)
            for raw_path in (image_paths or [])
        )
        candidates.extend(
            ("document", raw_path)
            for raw_path in (document_paths or [])
        )

        metadata: list[dict] = []
        seen_paths: set[str] = set()

        for kind, raw_path in candidates:
            source_path = Path(
                raw_path
            ).expanduser()

            if (
                not source_path.exists()
                or not source_path.is_file()
            ):
                continue

            normalized_source = str(
                source_path.resolve()
            ).lower()

            if normalized_source in seen_paths:
                continue

            seen_paths.add(
                normalized_source
            )

            safe_name = cls._safe_attachment_name(
                source_path.name
            )

            stored_name = (
                f"{uuid4().hex}_{safe_name}"
            )
            stored_path = (
                conversation_directory
                / stored_name
            )

            managed_copy = True

            try:
                shutil.copy2(
                    source_path,
                    stored_path,
                )
                final_path = stored_path

            except OSError:
                # Keep a non-managed reference as a fallback.
                managed_copy = False
                final_path = source_path.resolve()

            suffix = source_path.suffix.lower()

            metadata.append(
                {
                    "kind": kind,
                    "name": source_path.name,
                    "suffix": suffix,
                    "stored_path": str(
                        final_path
                    ),
                    "managed_copy": managed_copy,
                    "visual_memory": (
                        kind == "document"
                        and suffix == ".pdf"
                        and normalized_source
                        in visual_pdf_paths
                    ),
                }
            )

        return metadata

    @staticmethod
    def _remembered_pdf_image_urls(
        pdf_paths: list[str],
    ) -> list[str]:
        image_urls: list[str] = []

        for raw_path in pdf_paths[
            -MAX_REMEMBERED_VISUAL_PDFS:
        ]:
            path = Path(raw_path)

            if not path.exists() or not path.is_file():
                continue

            try:
                (
                    _text,
                    _notes,
                    rendered_image_urls,
                ) = ConversationManager._extract_pdf_content(
                    path
                )
            except Exception:
                continue

            remaining = (
                MAX_TOTAL_DOCUMENT_VISION_IMAGES
                - len(image_urls)
            )

            if remaining <= 0:
                break

            image_urls.extend(
                rendered_image_urls[:remaining]
            )

        return image_urls

    # ======================================================
    # Image conversion
    # ======================================================

    @staticmethod
    def _image_paths_to_data_urls(
        image_paths: list[str] | None,
    ) -> list[str]:
        if not image_paths:
            return []

        image_urls: list[str] = []

        for raw_path in image_paths:
            path = Path(raw_path).expanduser()

            if not path.exists() or not path.is_file():
                raise FileNotFoundError(
                    f"Attached image was not found: {path}"
                )

            suffix = path.suffix.lower()

            if suffix not in SUPPORTED_IMAGE_SUFFIXES:
                raise ValueError(
                    f"Unsupported image type: {path.name}"
                )

            mime_type = (
                mimetypes.guess_type(path.name)[0]
                or "image/jpeg"
            )

            encoded = base64.b64encode(
                path.read_bytes()
            ).decode("ascii")

            image_urls.append(
                f"data:{mime_type};base64,{encoded}"
            )

        return image_urls

    # ======================================================
    # Document extraction
    # ======================================================

    @staticmethod
    def _read_plain_text_document(
        path: Path,
    ) -> str:
        """Read a text/code file using common encodings."""

        encodings = (
            "utf-8",
            "utf-8-sig",
            "cp1252",
            "latin-1",
        )

        last_error: Exception | None = None

        for encoding in encodings:
            try:
                return path.read_text(
                    encoding=encoding,
                )
            except UnicodeDecodeError as exc:
                last_error = exc

        raise ValueError(
            f"Could not decode text from {path.name}: "
            f"{last_error}"
        )

    @staticmethod
    def _extract_pdf_content(
        path: Path,
    ) -> tuple[str, list[str], list[str]]:
        """
        Extract selectable PDF text and render low-text pages
        for the vision model.

        Returns:
            text: Selectable text grouped by page.
            notes: Extraction and page-limit notes.
            image_urls: Base64 data URLs for scanned pages.
        """

        import fitz
        from pypdf import PdfReader

        reader = PdfReader(str(path))

        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:
                unlocked = 0

            if not unlocked:
                raise ValueError(
                    f"{path.name} is password protected."
                )

        total_pages = len(reader.pages)
        page_limit = min(
            total_pages,
            MAX_PDF_PAGES,
        )

        extracted_pages: list[str] = []
        scanned_page_indexes: list[int] = []
        notes: list[str] = []

        for page_index in range(page_limit):
            page = reader.pages[page_index]

            try:
                page_text = (
                    page.extract_text() or ""
                ).strip()
            except Exception as exc:
                page_text = ""
                notes.append(
                    "Selectable-text extraction failed on "
                    f"page {page_index + 1}: {exc}"
                )

            compact_text = " ".join(
                page_text.split()
            )

            if page_text:
                extracted_pages.append(
                    f"[Page {page_index + 1}]\n"
                    f"{page_text}"
                )

            if (
                len(compact_text)
                < MIN_PDF_TEXT_CHARACTERS_PER_PAGE
            ):
                scanned_page_indexes.append(
                    page_index
                )

        if total_pages > page_limit:
            notes.append(
                f"Only the first {page_limit} of "
                f"{total_pages} PDF pages were processed."
            )

        image_urls: list[str] = []
        selected_scanned_pages = scanned_page_indexes[
            :MAX_SCANNED_PDF_VISION_PAGES
        ]

        if selected_scanned_pages:
            pdf_document = fitz.open(str(path))

            try:
                zoom = (
                    SCANNED_PDF_RENDER_DPI
                    / 72.0
                )
                matrix = fitz.Matrix(
                    zoom,
                    zoom,
                )

                for page_index in selected_scanned_pages:
                    rendered_page = (
                        pdf_document.load_page(
                            page_index
                        )
                    )

                    pixmap = rendered_page.get_pixmap(
                        matrix=matrix,
                        alpha=False,
                    )

                    try:
                        image_bytes = pixmap.tobytes(
                            "jpeg"
                        )
                        mime_type = "image/jpeg"
                    except Exception:
                        image_bytes = pixmap.tobytes(
                            "png"
                        )
                        mime_type = "image/png"

                    encoded = base64.b64encode(
                        image_bytes
                    ).decode("ascii")

                    image_urls.append(
                        f"data:{mime_type};base64,"
                        f"{encoded}"
                    )

            finally:
                pdf_document.close()

            rendered_numbers = [
                str(page_index + 1)
                for page_index
                in selected_scanned_pages
            ]

            notes.append(
                "Rendered low-text/scanned PDF page(s) "
                "for visual inspection: "
                + ", ".join(rendered_numbers)
                + "."
            )

        if (
            len(scanned_page_indexes)
            > len(selected_scanned_pages)
        ):
            omitted_count = (
                len(scanned_page_indexes)
                - len(selected_scanned_pages)
            )
            notes.append(
                f"{omitted_count} additional low-text/scanned "
                "PDF page(s) were not rendered because the "
                "vision-page limit was reached."
            )

        text = "\n\n".join(
            extracted_pages
        ).strip()

        if not text and image_urls:
            text = (
                "[No meaningful selectable text was found. "
                "The scanned PDF pages listed in the extraction "
                "notes were supplied to the vision model.]"
            )
        elif not text:
            text = (
                "[No selectable text or renderable scanned "
                "pages were found in this PDF.]"
            )

        return text, notes, image_urls

    @staticmethod
    def _extract_docx_text(
        path: Path,
    ) -> str:
        """Extract paragraphs and tables from a DOCX file."""

        from docx import Document

        document = Document(str(path))
        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):
            rows: list[str] = []

            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                rows.append(" | ".join(values))

            if rows:
                parts.append(
                    f"[Table {table_index}]\n"
                    + "\n".join(rows)
                )

        text = "\n\n".join(parts).strip()

        return text or (
            "[The DOCX file contained no extractable text.]"
        )

    @staticmethod
    def _clean_table_value(
        value,
    ) -> str:
        """Convert a spreadsheet/table cell into compact text."""

        if value is None:
            return ""

        text = str(value).replace(
            "\r\n",
            " ",
        ).replace(
            "\n",
            " ",
        ).replace(
            "\r",
            " ",
        )

        return " ".join(
            text.split()
        ).strip()

    @classmethod
    def _extract_csv_text(
        cls,
        path: Path,
    ) -> tuple[str, list[str]]:
        """Parse CSV data into a bounded, readable table."""

        import pandas as pd

        notes: list[str] = []
        frame = None
        last_error: Exception | None = None

        for encoding in (
            "utf-8",
            "utf-8-sig",
            "cp1252",
            "latin-1",
        ):
            try:
                frame = pd.read_csv(
                    path,
                    sep=None,
                    engine="python",
                    dtype=str,
                    keep_default_na=False,
                    encoding=encoding,
                    nrows=(
                        MAX_SPREADSHEET_ROWS_PER_SHEET
                        + 1
                    ),
                    on_bad_lines="skip",
                )
                break

            except UnicodeDecodeError as exc:
                last_error = exc

            except Exception as exc:
                last_error = exc

        if frame is None:
            raise ValueError(
                f"Could not parse CSV file {path.name}: "
                f"{last_error}"
            )

        if len(frame) > MAX_SPREADSHEET_ROWS_PER_SHEET:
            frame = frame.iloc[
                :MAX_SPREADSHEET_ROWS_PER_SHEET
            ]
            notes.append(
                "Only the first "
                f"{MAX_SPREADSHEET_ROWS_PER_SHEET} "
                "CSV data rows were processed."
            )

        if len(frame.columns) > MAX_SPREADSHEET_COLUMNS:
            frame = frame.iloc[
                :,
                :MAX_SPREADSHEET_COLUMNS,
            ]
            notes.append(
                "Only the first "
                f"{MAX_SPREADSHEET_COLUMNS} "
                "CSV columns were processed."
            )

        frame.columns = [
            cls._clean_table_value(column)
            or f"Column {index + 1}"
            for index, column in enumerate(
                frame.columns
            )
        ]

        for column in frame.columns:
            frame[column] = frame[column].map(
                cls._clean_table_value
            )

        table_text = frame.to_csv(
            sep="\t",
            index=False,
            lineterminator="\n",
        ).strip()

        return (
            "[CSV table]\n"
            + (
                table_text
                or "[The CSV file contained no data rows.]"
            ),
            notes,
        )

    @classmethod
    def _extract_excel_text(
        cls,
        path: Path,
    ) -> tuple[str, list[str]]:
        """Extract bounded cell values from XLSX/XLSM workbooks."""

        from openpyxl import load_workbook

        notes: list[str] = []
        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
        )

        try:
            sheet_names = workbook.sheetnames
            selected_names = sheet_names[
                :MAX_SPREADSHEET_SHEETS
            ]

            if len(sheet_names) > len(selected_names):
                notes.append(
                    "Only the first "
                    f"{MAX_SPREADSHEET_SHEETS} of "
                    f"{len(sheet_names)} workbook sheets "
                    "were processed."
                )

            sections: list[str] = []

            for sheet_name in selected_names:
                worksheet = workbook[sheet_name]
                rows: list[str] = []
                row_limit_hit = False
                column_limit_hit = False

                for row_number, row in enumerate(
                    worksheet.iter_rows(
                        values_only=True
                    ),
                    start=1,
                ):
                    if (
                        row_number
                        > MAX_SPREADSHEET_ROWS_PER_SHEET
                    ):
                        row_limit_hit = True
                        break

                    values = [
                        cls._clean_table_value(value)
                        for value in row[
                            :MAX_SPREADSHEET_COLUMNS
                        ]
                    ]

                    if len(row) > MAX_SPREADSHEET_COLUMNS:
                        column_limit_hit = True

                    while values and not values[-1]:
                        values.pop()

                    if not any(values):
                        continue

                    rows.append(
                        f"Row {row_number}: "
                        + " | ".join(values)
                    )

                if row_limit_hit:
                    notes.append(
                        f'Sheet "{sheet_name}" was limited '
                        "to the first "
                        f"{MAX_SPREADSHEET_ROWS_PER_SHEET} rows."
                    )

                if column_limit_hit:
                    notes.append(
                        f'Sheet "{sheet_name}" was limited '
                        "to the first "
                        f"{MAX_SPREADSHEET_COLUMNS} columns."
                    )

                sections.append(
                    f'[Sheet: {sheet_name}]\n'
                    + (
                        "\n".join(rows)
                        if rows
                        else "[No non-empty cell values found.]"
                    )
                )

            notes.append(
                "Formula cells use their last cached calculated "
                "values; uncached formulas may appear blank."
            )

            return (
                "\n\n".join(sections).strip()
                or "[The workbook contained no readable sheets.]",
                notes,
            )

        finally:
            workbook.close()

    @classmethod
    def _extract_pptx_text(
        cls,
        path: Path,
    ) -> tuple[str, list[str]]:
        """Extract visible text and table cells from PPTX slides."""

        from pptx import Presentation

        notes: list[str] = []
        presentation = Presentation(str(path))
        total_slides = len(presentation.slides)
        slide_limit = min(
            total_slides,
            MAX_PRESENTATION_SLIDES,
        )

        sections: list[str] = []
        visual_only_content_found = False

        for slide_number in range(slide_limit):
            slide = presentation.slides[
                slide_number
            ]
            parts: list[str] = []

            for shape in slide.shapes:
                if getattr(
                    shape,
                    "has_text_frame",
                    False,
                ):
                    text = cls._clean_table_value(
                        shape.text
                    )

                    if text:
                        parts.append(text)

                if getattr(
                    shape,
                    "has_table",
                    False,
                ):
                    table_rows: list[str] = []

                    for row in shape.table.rows:
                        values = [
                            cls._clean_table_value(
                                cell.text
                            )
                            for cell in row.cells
                        ]
                        table_rows.append(
                            " | ".join(values)
                        )

                    if table_rows:
                        parts.append(
                            "[Table]\n"
                            + "\n".join(table_rows)
                        )

                shape_type_name = str(
                    getattr(
                        shape,
                        "shape_type",
                        "",
                    )
                ).upper()

                if (
                    "PICTURE" in shape_type_name
                    or "CHART" in shape_type_name
                    or "MEDIA" in shape_type_name
                ):
                    visual_only_content_found = True

            sections.append(
                f"[Slide {slide_number + 1}]\n"
                + (
                    "\n".join(parts)
                    if parts
                    else "[No selectable text found.]"
                )
            )

        if total_slides > slide_limit:
            notes.append(
                f"Only the first {slide_limit} of "
                f"{total_slides} slides were processed."
            )

        if visual_only_content_found:
            notes.append(
                "Pictures, charts, diagrams, animations, and "
                "other visual-only slide content were not "
                "visually interpreted; selectable text and "
                "tables were extracted."
            )

        return (
            "\n\n".join(sections).strip()
            or "[The presentation contained no readable slides.]",
            notes,
        )

    def _extract_document_text(
        self,
        path: Path,
    ) -> tuple[str, list[str], list[str]]:
        """
        Extract text, limitation notes, and optional rendered
        document-page images from one file.
        """

        suffix = path.suffix.lower()
        notes: list[str] = []
        document_image_urls: list[str] = []

        if suffix == ".pdf":
            (
                text,
                pdf_notes,
                document_image_urls,
            ) = self._extract_pdf_content(
                path
            )
            notes.extend(pdf_notes)

        elif suffix == ".docx":
            text = self._extract_docx_text(
                path
            )

        elif suffix == ".csv":
            text, extraction_notes = (
                self._extract_csv_text(path)
            )
            notes.extend(extraction_notes)

        elif suffix in {".xlsx", ".xlsm"}:
            text, extraction_notes = (
                self._extract_excel_text(path)
            )
            notes.extend(extraction_notes)

        elif suffix == ".pptx":
            text, extraction_notes = (
                self._extract_pptx_text(path)
            )
            notes.extend(extraction_notes)

        elif suffix in PLAIN_TEXT_DOCUMENT_SUFFIXES:
            text = self._read_plain_text_document(
                path
            )

        else:
            raise ValueError(
                f"Unsupported document type: {path.name}"
            )

        text = text.strip()

        if len(text) > MAX_DOCUMENT_CHARACTERS_PER_FILE:
            text = text[
                :MAX_DOCUMENT_CHARACTERS_PER_FILE
            ].rstrip()
            notes.append(
                "The extracted text was shortened to "
                f"{MAX_DOCUMENT_CHARACTERS_PER_FILE:,} "
                "characters for this request."
            )

        return (
            text,
            notes,
            document_image_urls,
        )

    def _documents_to_context(
        self,
        document_paths: list[str] | None,
    ) -> tuple[str, list[str], set[str]]:
        """
        Build bounded text context and rendered scanned-PDF
        page images from local documents.
        """

        if not document_paths:
            return "", [], set()

        sections: list[str] = []
        document_image_urls: list[str] = []
        visual_pdf_paths: set[str] = set()
        total_characters = 0

        for raw_path in document_paths:
            path = Path(raw_path).expanduser()

            if not path.exists() or not path.is_file():
                raise FileNotFoundError(
                    f"Attached document was not found: {path}"
                )

            if path.suffix.lower() not in SUPPORTED_DOCUMENT_SUFFIXES:
                raise ValueError(
                    f"Unsupported document type: {path.name}"
                )

            remaining = (
                MAX_DOCUMENT_CHARACTERS_TOTAL
                - total_characters
            )

            if remaining <= 0:
                sections.append(
                    "[Additional documents were omitted because "
                    "the combined extraction limit was reached.]"
                )
                break

            (
                text,
                notes,
                extracted_image_urls,
            ) = self._extract_document_text(
                path
            )

            if (
                path.suffix.lower() == ".pdf"
                and extracted_image_urls
            ):
                visual_pdf_paths.add(
                    str(path.resolve())
                )

            remaining_image_slots = (
                MAX_TOTAL_DOCUMENT_VISION_IMAGES
                - len(document_image_urls)
            )

            if remaining_image_slots > 0:
                document_image_urls.extend(
                    extracted_image_urls[
                        :remaining_image_slots
                    ]
                )

            if (
                len(extracted_image_urls)
                > max(0, remaining_image_slots)
            ):
                omitted_images = (
                    len(extracted_image_urls)
                    - max(0, remaining_image_slots)
                )
                notes.append(
                    f"{omitted_images} rendered document "
                    "page image(s) were omitted because the "
                    "combined vision-image limit was reached."
                )

            if len(text) > remaining:
                text = text[:remaining].rstrip()
                notes.append(
                    "This file was further shortened because "
                    "the combined document limit was reached."
                )

            total_characters += len(text)

            note_block = ""

            if notes:
                note_block = (
                    "\n\nExtraction notes:\n- "
                    + "\n- ".join(notes)
                )

            sections.append(
                f"<document name={path.name!r}>\n"
                f"{text}"
                f"{note_block}\n"
                "</document>"
            )

        if not sections:
            return (
                "",
                document_image_urls,
                visual_pdf_paths,
            )

        context = (
            "The following local document content was "
            "extracted for this request. Treat it as user-"
            "provided source material.\n\n"
            + "\n\n".join(sections)
        )

        return (
            context,
            document_image_urls,
            visual_pdf_paths,
        )

    # ======================================================
    # Conversation title
    # ======================================================

    @staticmethod
    def _title_from_prompt(
        prompt: str,
    ) -> str:
        one_line = " ".join(
            prompt.split()
        ).strip()

        if not one_line:
            return "Image Chat"

        maximum_length = 45

        if len(one_line) <= maximum_length:
            return one_line

        return (
            one_line[:maximum_length].rstrip()
            + "..."
        )

    # ======================================================
    # Ask Keshav
    # ======================================================

    async def ask(
        self,
        prompt: str,
        conversation_id: str | None = None,
        image_paths: list[str] | None = None,
        document_paths: list[str] | None = None,
        on_chunk: Callable[[str], None] | None = None,
        live_mode: bool = False,
    ) -> str:
        """
        Send the conversation directly to native Ollama.

        When on_chunk is provided, each generated text chunk is
        delivered immediately while the complete response is still
        accumulated and stored normally.
        """

        clean_prompt = prompt.strip()

        resolved_id = self._resolve_conversation_id(
            conversation_id,
        )

        session = self._get_session(
            resolved_id,
        )

        image_urls = self._image_paths_to_data_urls(
            image_paths
        )

        (
            document_context,
            document_image_urls,
            visual_pdf_paths,
        ) = await asyncio.to_thread(
            self._documents_to_context,
            document_paths,
        )

        remembered_image_urls: list[str] = []
        remembered_pdf_image_urls: list[str] = []

        # Re-send recent managed visual attachments only when
        # the user is asking a follow-up without new files.
        if not image_paths and not document_paths:
            with self._sessions_lock:
                remembered_image_paths = list(
                    session.remembered_image_paths
                )
                remembered_visual_pdf_paths = list(
                    session.remembered_visual_pdf_paths
                )

            remembered_image_urls = (
                self._image_paths_to_data_urls(
                    remembered_image_paths
                )
            )

            remembered_pdf_image_urls = (
                await asyncio.to_thread(
                    self._remembered_pdf_image_urls,
                    remembered_visual_pdf_paths,
                )
            )

        all_image_urls = (
            image_urls
            + document_image_urls
            + remembered_image_urls
            + remembered_pdf_image_urls
        )

        if (
            not clean_prompt
            and all_image_urls
            and document_context
        ):
            clean_prompt = (
                "Please analyze the attached images and "
                "documents together."
            )
        elif not clean_prompt and all_image_urls:
            clean_prompt = (
                "Please describe and analyze the attached "
                "image or images."
            )
        elif not clean_prompt and document_context:
            clean_prompt = (
                "Please summarize and analyze the attached "
                "document or documents."
            )

        if not clean_prompt:
            return "Please enter a message."

        model_prompt = clean_prompt

        if document_context:
            model_prompt = (
                clean_prompt
                + "\n\n--- ATTACHED DOCUMENT CONTENT ---\n"
                + document_context
                + "\n--- END ATTACHED DOCUMENT CONTENT ---"
            )

        if (
            remembered_image_urls
            or remembered_pdf_image_urls
        ):
            model_prompt = (
                model_prompt
                + "\n\n[Previously attached visual files "
                "are included again for follow-up context.]"
            )

        await asyncio.to_thread(
            session.request_lock.acquire
        )

        user_message = Message.user_message(
            model_prompt,
        )

        try:
            with self._sessions_lock:
                is_first_message = (
                    len(session.messages) == 0
                )

                session.messages.append(
                    user_message
                )

                self._trim_history(
                    session
                )

                history_snapshot = list(
                    session.messages
                )

                if live_mode:
                    history_snapshot = history_snapshot[
                        -LIVE_HISTORY_MESSAGES:
                    ]

            try:
                selected_client = (
                    self.vision_llm
                    if all_image_urls
                    else self.llm
                )

                response_parts: list[str] = []

                active_system_prompt = SYSTEM_PROMPT

                if live_mode:
                    active_system_prompt = (
                        SYSTEM_PROMPT
                        + "\n"
                        + LIVE_MODE_PROMPT
                    )

                async for chunk in selected_client.stream_chat(
                    messages=history_snapshot,
                    system_prompt=active_system_prompt,
                    images=(
                        all_image_urls
                        if all_image_urls
                        else None
                    ),
                    fast_mode=live_mode,
                ):
                    response_parts.append(
                        chunk
                    )

                    if on_chunk is not None:
                        try:
                            on_chunk(chunk)
                        except Exception:
                            # UI streaming must never break the
                            # actual model request.
                            pass

                response = "".join(
                    response_parts
                )

            except Exception:
                # Do not keep failed user requests in memory.
                with self._sessions_lock:
                    if (
                        session.messages
                        and session.messages[-1]
                        is user_message
                    ):
                        session.messages.pop()
                    else:
                        try:
                            session.messages.remove(
                                user_message
                            )
                        except ValueError:
                            pass

                raise

            response = (
                response.strip()
                if response
                else "I couldn't generate a response."
            )

            assistant_message = (
                Message.assistant_message(
                    response
                )
            )

            with self._sessions_lock:
                session.messages.append(
                    assistant_message
                )
                self._trim_history(
                    session
                )

            stored_attachments = (
                await asyncio.to_thread(
                    self._persist_attachment_files,
                    resolved_id,
                    image_paths,
                    document_paths,
                    visual_pdf_paths,
                )
            )

            chat_store.add_message(
                conversation_id=resolved_id,
                role="user",
                content=clean_prompt,
                model_content=model_prompt,
                attachments=stored_attachments,
            )

            chat_store.add_message(
                conversation_id=resolved_id,
                role="assistant",
                content=response,
                model_content=response,
            )

            new_remembered_images = [
                attachment["stored_path"]
                for attachment in stored_attachments
                if (
                    attachment.get("kind") == "image"
                    and attachment.get("stored_path")
                )
            ]

            new_visual_pdfs = [
                attachment["stored_path"]
                for attachment in stored_attachments
                if (
                    attachment.get("kind") == "document"
                    and attachment.get("visual_memory")
                    and attachment.get("stored_path")
                )
            ]

            with self._sessions_lock:
                session.remembered_image_paths.extend(
                    new_remembered_images
                )
                session.remembered_image_paths = (
                    session.remembered_image_paths[
                        -MAX_REMEMBERED_IMAGE_ATTACHMENTS:
                    ]
                )

                session.remembered_visual_pdf_paths.extend(
                    new_visual_pdfs
                )
                session.remembered_visual_pdf_paths = (
                    session.remembered_visual_pdf_paths[
                        -MAX_REMEMBERED_VISUAL_PDFS:
                    ]
                )

            if is_first_message:
                chat_store.rename_conversation(
                    conversation_id=resolved_id,
                    title=self._title_from_prompt(
                        clean_prompt
                    ),
                )

            return response

        finally:
            session.request_lock.release()

    # ======================================================
    # Create conversations
    # ======================================================

    def create_conversation(
        self,
        title: str = "New Chat",
        make_active: bool = True,
    ) -> str:
        conversation_id = (
            chat_store.create_conversation(
                title=title,
            )
        )

        with self._sessions_lock:
            self._sessions[
                conversation_id
            ] = ConversationSession()

        if make_active:
            self._set_active_conversation(
                conversation_id
            )

        return conversation_id

    def start_new_conversation(self) -> str:
        chat_store.delete_empty_conversations()
        self._remove_stale_sessions()

        return self.create_conversation(
            title="New Chat",
            make_active=True,
        )

    # ======================================================
    # Load/read state
    # ======================================================

    def load_conversation(
        self,
        conversation_id: str,
    ) -> list[Message]:
        if not chat_store.conversation_exists(
            conversation_id,
        ):
            raise ValueError(
                "Conversation does not exist."
            )

        session = self._get_session(
            conversation_id,
        )

        self._set_active_conversation(
            conversation_id,
        )

        with self._sessions_lock:
            return list(
                session.messages
            )

    def get_active_conversation_id(self) -> str:
        with self._sessions_lock:
            return self._active_conversation_id

    def get_history(
        self,
        conversation_id: str | None = None,
    ) -> list[Message]:
        resolved_id = self._resolve_conversation_id(
            conversation_id,
        )

        session = self._get_session(
            resolved_id,
        )

        with self._sessions_lock:
            return list(
                session.messages
            )

    def list_conversations(self) -> list[dict]:
        conversations = chat_store.list_conversations(
            include_empty=True,
        )

        active_id = self.get_active_conversation_id()

        return [
            conversation
            for conversation in conversations
            if (
                conversation["message_count"] > 0
                or conversation["id"] == active_id
            )
        ]

    # ======================================================
    # Rename/delete
    # ======================================================

    def rename_conversation(
        self,
        conversation_id: str,
        title: str,
    ) -> None:
        chat_store.rename_conversation(
            conversation_id=conversation_id,
            title=title,
        )

    def delete_conversation(
        self,
        conversation_id: str,
    ) -> str:
        chat_store.delete_conversation(
            conversation_id,
        )

        with self._sessions_lock:
            self._sessions.pop(
                conversation_id,
                None,
            )

            active_was_deleted = (
                conversation_id
                == self._active_conversation_id
            )

        if active_was_deleted:
            conversations = (
                chat_store.list_conversations(
                    include_empty=False,
                )
            )

            if conversations:
                next_id = conversations[0]["id"]
                self._set_active_conversation(
                    next_id
                )
            else:
                next_id = self.create_conversation(
                    title="New Chat",
                    make_active=True,
                )

            return next_id

        return self.get_active_conversation_id()


conversation_manager = ConversationManager()


async def ask_keshav(
    prompt: str,
    conversation_id: str | None = None,
    image_paths: list[str] | None = None,
    document_paths: list[str] | None = None,
    on_chunk: Callable[[str], None] | None = None,
    live_mode: bool = False,
) -> str:
    """
    Ask Keshav directly through Ollama.

    on_chunk receives streamed response text immediately.
    """

    return await conversation_manager.ask(
        prompt=prompt,
        conversation_id=conversation_id,
        image_paths=image_paths,
        document_paths=document_paths,
        on_chunk=on_chunk,
        live_mode=live_mode,
    )


def clear_conversation(
    conversation_id: str | None = None,
) -> str:
    del conversation_id

    return (
        conversation_manager
        .start_new_conversation()
    )


def create_conversation(
    title: str = "New Chat",
) -> str:
    return (
        conversation_manager
        .create_conversation(
            title=title,
            make_active=True,
        )
    )


def load_conversation(
    conversation_id: str,
) -> list[Message]:
    return (
        conversation_manager
        .load_conversation(
            conversation_id
        )
    )


def delete_conversation(
    conversation_id: str,
) -> str:
    return (
        conversation_manager
        .delete_conversation(
            conversation_id
        )
    )


def rename_conversation(
    conversation_id: str,
    title: str,
) -> None:
    conversation_manager.rename_conversation(
        conversation_id=conversation_id,
        title=title,
    )


def list_conversations() -> list[dict]:
    return conversation_manager.list_conversations()


def get_active_conversation_id() -> str:
    return (
        conversation_manager
        .get_active_conversation_id()
    )


def get_conversation_history(
    conversation_id: str | None = None,
) -> list[Message]:
    return conversation_manager.get_history(
        conversation_id
    )
